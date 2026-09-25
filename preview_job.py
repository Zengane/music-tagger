"""Makes what the preview needs from files the window cannot open itself.

    python preview_job.py mesh  <model>    <out.glb>
    python preview_job.py pages <document> <out folder> [max pages] [width]
    python preview_job.py text  <document> <out.txt>

The window draws PNGs and loads glTF, OBJ and VOX; everything else comes
through here once and is cached (see ui/preview.odin):

    mesh    any 3D format this app knows -> GLB. trimesh for STL, PLY, OFF,
            3MF and DAE; the converters in mtag/media for Blockbench and
            MagicaVoxel; Blender, when installed, for FBX, .blend, USD,
            Alembic, 3DS and X3D.
    pages   PDF pages as PNGs, with pypdfium2 (Apache-2.0 / BSD-3).
    text    the words of a .docx, .odt or .rtf, for reading in the viewer.

Prints one JSON line: {"ok": true, ...} or {"ok": false, "error": "..."}.
"""

from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from mtag import console


def mesh(src: Path, out: Path) -> dict:
    from mtag.media.render3d import NEEDS_BLENDER, NEEDS_OWN_APP, _via_blender

    suffix = src.suffix.lower()
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.glb")
    if suffix in NEEDS_OWN_APP:
        return {"ok": False, "needs_app": NEEDS_OWN_APP[suffix],
                "error": f"{src.suffix} files open only in {NEEDS_OWN_APP[suffix]}. "
                         f"Export it as FBX, OBJ or glTF to preview it here."}
    if suffix == ".bbmodel":
        from mtag.media.bbmodel import to_glb

        to_glb(src, tmp)
    elif suffix == ".vox":
        from mtag.media.voxel import to_glb

        to_glb(src, tmp)
    elif suffix in NEEDS_BLENDER:
        converted = _via_blender(src)
        tmp.write_bytes(converted.read_bytes())
    else:
        import trimesh

        scene = trimesh.load(str(src), force="scene")
        if not scene.geometry:
            return {"ok": False, "error": f"no geometry in {src.name}"}
        tmp.write_bytes(scene.export(file_type="glb"))
    tmp.replace(out)
    return {"ok": True, "out": str(out)}


def pages(src: Path, out_dir: Path, max_pages: int = 40, width: int = 1400) -> dict:
    import pypdfium2 as pdfium

    out_dir.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(src))
    total = len(pdf)
    written = 0
    for i in range(min(total, max_pages)):
        page = pdf[i]
        w, _h = page.get_size()
        image = page.render(scale=max(width / max(w, 1), 0.1)).to_pil()
        image.convert("RGB").save(out_dir / f"page-{i + 1:03d}.png")
        written += 1
    (out_dir / "pages.json").write_text(json.dumps({"total": total, "written": written}), encoding="utf-8")
    return {"ok": True, "total": total, "written": written}


def _docx_text(src: Path) -> str:
    import docx

    d = docx.Document(str(src))
    return "\n\n".join(p.text for p in d.paragraphs)


def _odt_text(src: Path) -> str:
    with zipfile.ZipFile(src) as z:
        xml = z.read("content.xml").decode("utf-8", "replace")
    xml = re.sub(r"</text:p>|<text:line-break/>|</text:h>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def _rtf_text(src: Path) -> str:
    raw = src.read_text(encoding="latin-1", errors="replace")
    raw = re.sub(r"\\par[d]?", "\n", raw)
    raw = re.sub(r"\{\\\*[^{}]*\}", "", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
    return raw.replace("{", "").replace("}", "")


def text(src: Path, out: Path) -> dict:
    suffix = src.suffix.lower()
    if suffix == ".docx":
        body = _docx_text(src)
    elif suffix == ".odt":
        body = _odt_text(src)
    elif suffix == ".rtf":
        body = _rtf_text(src)
    else:
        body = src.read_text(encoding="utf-8", errors="replace")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(body.strip() + "\n", encoding="utf-8")
    return {"ok": True, "chars": len(body)}


def main(argv: list[str]) -> int:
    console.init()
    if len(argv) < 3:
        print(json.dumps({"ok": False, "error": "usage: preview_job.py mesh|pages|text <src> <out>"}))
        return 2
    what, src, out = argv[0], Path(argv[1]), Path(argv[2])
    try:
        if what == "mesh":
            result = mesh(src, out)
        elif what == "pages":
            result = pages(src, out, *(int(a) for a in argv[3:5]))
        elif what == "text":
            result = text(src, out)
        else:
            result = {"ok": False, "error": f"unknown preview {what}"}
    except Exception as e:  # reported, not raised: the window shows the reason
        result = {"ok": False, "error": f"{type(e).__name__}: {e}"}
    print(json.dumps(result), flush=True)
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
